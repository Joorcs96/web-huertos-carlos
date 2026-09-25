import openpyxl
import json
import datetime
import re
import os

wb = openpyxl.load_workbook('tools/huertos_carlos.xlsx', data_only=True)

PLAGAS_PATTERNS = {
    'Araña roja': [
        r'\baraña\b', r'\baranya\b', r'\bàcar\b', r'\bacaro\b', r'\bácaro\b', r'\btetranychus\b',
        r'\bar\b'
    ],
    'Mosca de la fruta': [
        r'\bceratitis\b',
        r'\btrampes?\s+mosca\b',
        r'\bmaxina\s+mosca\b',
        r'\bmosca\s+de\s+la\s+fruita\b',
        r'\bmosca\s+de\s+la\s+fruta\b',
        r'\bcuidar\s+mosca\b',
        r'\bcontrol\s+ceratitis\b',
        r'\bmosca\b'
    ],
    'Mosca blanca': [
        r'\bmosca\s+blanca\b',
        r'\baleurothrixus\b',
        r'\bmb\b'
    ],
    'Trip': [
        r'\btrips?\b', r'\btryps?\b', r'\bpeiró\b', r'\bpeiro\b', r'\bscirtothrips\b'
    ],
    'Cotonet': [
        r'\bcotonet\b', r'\bplanococcus\b', r'\bpseudococcus\b'
    ],
    'Piojo rojo': [
        r'\bpiojo\s+rojo\b', r'\bpoll\s+roig\b', r'\bprc\b', r'\bserpeta\b', r'\bcaparreta\b',
        r'\bpiojo\b(?!.*blanco)'
    ],
    'Minador': [
        r'\bminador\b', r'\bminaor\b', r'\bphyllocnistis\b'
    ]
}

NEGATION_PATTERNS = [
    r'\bno\s+(?:hi\s+ha|hia|hi\s+han|n[\'’]hi\s+ha|es\s+veu|se\s+veu|veu|te|té|tenim|queden|arriba|sintomes?|danys?|ha\s+tingut|ha\s+tingut\s+mai|ha\s+tenido|veig|senyals?|señals?)\b',
    r'\bsense\s+(?:mal\s+de|massa\s+mal\s+de|masa\s+mal\s+de|gens\s+de|presencia\s+de|danys?\s+de)?\b',
    r'\b(?:limpio|limpia|net|neta|res|nada)\s+de\b',
    r'\bapenas\s+se\s+ven?\b',
    r'\bno\s+(?:hi\s+ha\s+)?plagues\b',
    r'\bsense\s+plagues\b',
    r'\bno\s+plsgues\b',
    r'\bno\s+te\s+ni\b',
    r'\bno\s+hi\s+ha\s+cap\b'
]

def is_negated_clause(clause, span):
    pre = clause[:span[0]].lower()
    post = clause[span[1]:].lower()
    
    # 1. Comprobar si inmediatamente antes hay 'no', 'ni', 'sense', 'res de', etc.
    if re.search(r'\b(?:no|ni|sense|net|neta|limpio|limpia|zero|cero|res|nada)\s+(?:gaire|massa|masa|gens|pas|de)?\s*$', pre):
        return True
        
    # 2. Comprobar patrones de negación en la misma cláusula
    for neg in NEGATION_PATTERNS:
        neg_m = re.search(neg, pre)
        if neg_m:
            between = pre[neg_m.end():].strip()
            if not re.search(r'\b(?:pero|però|encara|tot\s+i)\b', between):
                return True
                
    # 3. Comprobar si inmediatamente después indica ausencia/cero o que está limpio
    if re.search(r'^\s*(?:cero|zero|inexistent|no\s+arriba|limpia|limpio|controlada?|casi\s+cero|parada|parat)', post):
        return True
        
    return False

def extract_plagas_and_negations(texto, faena_nom=""):
    comb = f"{faena_nom} {texto}".strip()
    if not comb:
        return [], []
        
    comb_lower = comb.lower()
    
    # Faenas directas de tratamiento fitosanitario
    if 'turbo ar' in comb_lower:
        return ['Araña roja'], []
    if 'turbo cotonet' in comb_lower:
        return ['Cotonet'], []
    if 'trampes mosca' in comb_lower or 'maxina mosca' in comb_lower:
        return ['Mosca de la fruta'], []

    clauses = re.split(r'[.;\n]+', comb_lower)
    plagas_activas = set()
    plagas_negadas = set()
    
    for cl in clauses:
        cl_clean = cl.strip()
        if not cl_clean:
            continue
            
        if re.search(r'\b(?:no|sense)\s+plagues\b', cl_clean):
            continue
            
        for plaga, patterns in PLAGAS_PATTERNS.items():
            for pat in patterns:
                for match in re.finditer(pat, cl_clean):
                    if is_negated_clause(cl_clean, match.span()):
                        plagas_negadas.add(plaga)
                    else:
                        plagas_activas.add(plaga)
                        
    # Desambiguación entre Mosca blanca y Mosca de la fruta
    if 'Mosca blanca' in plagas_activas and 'Mosca de la fruta' in plagas_activas:
        if not re.search(r'\b(?:ceratitis|trampes?|maxina|fruita|fruta)\b', comb_lower):
            plagas_activas.discard('Mosca de la fruta')
            
    # Si está activa y a la vez negada, predomina la presencia activa en el huerto
    plagas_negadas = plagas_negadas - plagas_activas
    
    return sorted(list(plagas_activas)), sorted(list(plagas_negadas))

def serialize_val(val):
    if val is None:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, float):
        return round(val, 4)
    return str(val).strip()

parcelas = []
faenas = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    max_r = ws.max_row
    max_c = ws.max_column
    
    info_col = None
    for r in range(1, min(10, max_r + 1)):
        for c in range(1, max_c + 1):
            val = ws.cell(r, c).value
            if val and str(val).strip().upper() == 'INFO':
                info_col = c
                break
        if info_col: break
        
    p_data = {
        'id': 'p-' + sheet_name.lower().replace(' ', '-').replace('.', '').replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u'),
        'nombre': sheet_name,
        'cultiu': '',
        'varietat': '',
        'patron': '',
        'marco': '',
        'superficieHa': '',
        'superficieFa': '',
        'codic': '',
        'anyPlantacio': '',
        'subparcelas': []
    }
    
    if info_col:
        for r in range(1, min(40, max_r + 1)):
            lbl = ws.cell(r, info_col).value
            v = ws.cell(r, info_col + 1).value
            if lbl:
                lbl_s = str(lbl).strip().lower()
                v_s = serialize_val(v)
                if 'cultiu' in lbl_s and not p_data['cultiu']: p_data['cultiu'] = v_s
                elif 'variet' in lbl_s and not p_data['varietat']: p_data['varietat'] = v_s
                elif 'patr' in lbl_s and not p_data['patron']: p_data['patron'] = v_s
                elif 'marc' in lbl_s and not p_data['marco']: p_data['marco'] = v_s
                elif 'finca (fa)' in lbl_s or 'superficie (fa)' in lbl_s: p_data['superficieFa'] = f"{v_s} Hanegadas" if v_s else ""
                elif 'finca (ha)' in lbl_s or 'superficie total' in lbl_s:
                    if not p_data['superficieHa']: p_data['superficieHa'] = f"{v_s} Ha" if v_s else ""
                elif 'codic' in lbl_s and not p_data['codic']: p_data['codic'] = v_s
                elif 'any de plantaci' in lbl_s and not p_data['anyPlantacio']: p_data['anyPlantacio'] = v_s
                elif 'parcela' in lbl_s and v_s and v_s not in p_data['subparcelas']: p_data['subparcelas'].append(v_s)
                
    sup_txt = []
    if p_data['superficieFa']: sup_txt.append(str(p_data['superficieFa']))
    if p_data['superficieHa']: sup_txt.append(f"({p_data['superficieHa']})")
    p_data['superficie'] = " ".join(sup_txt) if sup_txt else "Superficie en ficha"
    parcelas.append(p_data)
    
    # 2. Extraer FAENAS (Columnas A - D)
    for r in range(2, max_r + 1):
        fetxa = ws.cell(r, 1).value
        faena_nom = ws.cell(r, 2).value
        persona = ws.cell(r, 3).value
        comentari = ws.cell(r, 4).value
        
        if fetxa or faena_nom or comentari:
            f_str = serialize_val(fetxa)
            if not f_str or f_str.lower() == 'fetxa':
                continue
                
            com_str = serialize_val(comentari)
            faena_str = serialize_val(faena_nom)
            
            quimico = ""
            if any(k in faena_str.lower() or k in com_str.lower() for k in ['turbo', 't1', 't2', 't5', 't9', 't10', 't11', 't15', 't16', 't21', 't23', 't27', 't28', 't31', 'herbicida', 'abamectina', 'fe', 'abon', 'spintor', 'cobre', 'trebon', 'trampes']):
                quimico = f"{faena_str} - {com_str}" if com_str else faena_str

            plagas_act, plagas_neg = extract_plagas_and_negations(com_str, faena_str)
            es_tratamiento = bool('turbo' in faena_str.lower() or 'trampes' in faena_str.lower() or 'maxina' in faena_str.lower())
            
            faenas.append({
                'id': f"faena-{sheet_name.lower().replace(' ', '-')}-{r}",
                'parcelaId': p_data['id'],
                'parcelaNombre': sheet_name,
                'fecha': f_str,
                'hora': '10:00',
                'tipoFaena': faena_str if faena_str else 'Trabajos de campo',
                'usuario': serialize_val(persona) if persona else 'Equipo',
                'usuarioId': 'carlos' if 'carlos' in str(persona).lower() else 'operario1',
                'quimicoProducto': quimico,
                'quimicoDosis': com_str if quimico else '',
                'plagas': plagas_act,
                'plagasNegadas': plagas_neg,
                'esTratamiento': es_tratamiento,
                'hierba': 'Poca hierba' if 'herbicida' in (faena_str + ' ' + com_str).lower() else 'Limpio',
                'notas': com_str
            })

    # 3. Extraer INFORMES (Columnas F - K)
    for r in range(2, max_r + 1):
        inf_texto = None
        texto_cand = ws.cell(r, 6).value
        if texto_cand and isinstance(texto_cand, str) and len(texto_cand) > 12:
            inf_texto = texto_cand
            
        if inf_texto:
            fecha_encontrada = ""
            persona_encontrada = "Carlos"
            for back_r in range(r, max(1, r - 5), -1):
                for c in range(5, 12):
                    if ws.cell(back_r, c).value and str(ws.cell(back_r, c).value).strip().lower() == 'fetxa':
                        fecha_encontrada = serialize_val(ws.cell(back_r, c + 1).value)
                    if ws.cell(back_r, c).value and str(ws.cell(back_r, c).value).strip().lower() == 'persona':
                        persona_encontrada = serialize_val(ws.cell(back_r, c + 1).value)
                        
            plagas_inf_act, plagas_inf_neg = extract_plagas_and_negations(inf_texto, 'Informe')
            
            hierba_estado = 'Limpio'
            txt_lower = inf_texto.lower()
            if 'molta brossa' in txt_lower or 'mucha hierba' in txt_lower or 'plena de brossa' in txt_lower:
                hierba_estado = 'Mucha hierba'
            elif 'brossa' in txt_lower or 'hierba' in txt_lower:
                hierba_estado = 'Poca hierba'
                
            faenas.append({
                'id': f"informe-{sheet_name.lower().replace(' ', '-')}-{r}",
                'parcelaId': p_data['id'],
                'parcelaNombre': sheet_name,
                'fecha': fecha_encontrada if fecha_encontrada else '2026-07-01',
                'hora': '12:00',
                'tipoFaena': 'Informe de Estado / Revisión',
                'usuario': persona_encontrada if persona_encontrada else 'Carlos',
                'usuarioId': 'carlos',
                'quimicoProducto': '',
                'quimicoDosis': '',
                'plagas': plagas_inf_act,
                'plagasNegadas': plagas_inf_neg,
                'esTratamiento': False,
                'hierba': hierba_estado,
                'notas': inf_texto
            })

print(f"Total parcelas extraidas: {len(parcelas)}")
print(f"Total faenas/informes extraidos: {len(faenas)}")

with open('tools/datos_huertos_extraidos.json', 'w', encoding='utf-8') as f:
    json.dump({'parcelas': parcelas, 'faenas': faenas}, f, ensure_ascii=False, indent=2)

js_content = f"window.DATOS_INICIALES_CARLOS = {json.dumps({'parcelas': parcelas, 'faenas': faenas}, ensure_ascii=False, indent=2)};\n"
with open('datos_huertos.js', 'w', encoding='utf-8') as f:
    f.write(js_content)

print("Datos procesados y actualizados con éxito en datos_huertos.js y tools/datos_huertos_extraidos.json.")
